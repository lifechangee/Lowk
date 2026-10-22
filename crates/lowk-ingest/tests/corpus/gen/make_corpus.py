#!/usr/bin/env python3
"""Generates the RU/EN legal corpus of crates/lowk-ingest/tests/corpus (outputs are committed).
All text is invented for this repository. Needs: reportlab python-docx python-pptx openpyxl pillow and a
Cyrillic TTF (macOS: /Library/Fonts/Arial Unicode.ttf).
  uv run --with reportlab --with python-docx --with python-pptx --with openpyxl --with pillow gen/make_corpus.py .
Also writes truth/<file>.txt (the plain text a perfect extraction would contain) and MANIFEST.tsv."""
import sys, pathlib, textwrap
from reportlab.lib.pagesizes import A4
from reportlab.pdfgen import canvas
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from docx import Document
from pptx import Presentation
from pptx.util import Inches
import openpyxl
from PIL import Image, ImageDraw, ImageFont

out = pathlib.Path(sys.argv[1]).resolve()
(out / "truth").mkdir(exist_ok=True)
FONT = next(p for p in ["/Library/Fonts/Arial Unicode.ttf", "/System/Library/Fonts/Supplemental/Arial.ttf",
                        "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"] if pathlib.Path(p).exists())
pdfmetrics.registerFont(TTFont("F", FONT))
manifest = []

RU_PAGES = [
    ["ДОГОВОР ПОСТАВКИ № 17/2026",
     "г. Москва, 14 марта 2026 года",
     "Общество с ограниченной ответственностью «Северный ветер» (далее Поставщик) и Акционерное общество «Балтийская логистика» (далее Покупатель) заключили настоящий договор о нижеследующем.",
     "1. Предмет договора. Поставщик обязуется передать в собственность Покупателя оборудование согласно спецификации, а Покупатель обязуется принять товар и оплатить его.",
     "2. Срок поставки. Товар поставляется в течение тридцати календарных дней с момента получения предоплаты. Просрочка поставки влечёт неустойку в размере 0,1 процента от стоимости товара за каждый день."],
    ["3. Цена и порядок расчётов. Общая стоимость товара составляет 2 480 000 рублей, включая НДС 20 процентов. Покупатель вносит предоплату в размере 30 процентов в течение пяти рабочих дней.",
     "4. Качество и гарантия. Гарантийный срок составляет двадцать четыре месяца со дня подписания акта приёмки. Дефекты устраняются Поставщиком за его счёт в течение четырнадцати дней.",
     "5. Конфиденциальность. Стороны обязуются не разглашать условия договора третьим лицам без письменного согласия другой стороны."],
    ["6. Ответственность сторон. За неисполнение обязательств стороны несут ответственность в соответствии с законодательством Российской Федерации.",
     "7. Разрешение споров. Споры разрешаются путём переговоров, а при недостижении согласия в Арбитражном суде города Москвы.",
     "8. Заключительные положения. Договор вступает в силу с момента подписания и действует до полного исполнения обязательств. Подписи сторон: Иванов А. П., Петрова Е. С."],
]
EN_PAGES = [
    ["MUTUAL NON-DISCLOSURE AGREEMENT",
     "This Agreement is made on 2 April 2026 between Harbor Analytics Ltd, a company registered in England (the Discloser), and Lumen Robotics Inc, a Delaware corporation (the Recipient).",
     "1. Confidential Information means any non-public technical, financial or commercial information disclosed in writing or orally, including source code, pricing models and customer lists.",
     "2. Obligations. The Recipient shall keep the Confidential Information strictly confidential, use it solely to evaluate the proposed partnership, and protect it with at least reasonable care."],
    ["3. Exclusions. The obligations do not apply to information that is public through no fault of the Recipient, was already known to it, or is independently developed.",
     "4. Term. This Agreement remains in force for three years. The duty of confidentiality survives termination for a further period of five years.",
     "5. Governing law. This Agreement is governed by the laws of England and Wales, and the courts of London have exclusive jurisdiction. Signed by Mary Cooper and Daniel Reyes."],
]
MIX_PAGES = [
    ["Соглашение об уровне услуг / Service Level Agreement (SLA)",
     "Исполнитель (Provider) обеспечивает доступность сервиса не менее 99.9 процента в месяц. The Provider guarantees monthly availability of at least 99.9 percent.",
     "Инциденты уровня P1 принимаются в работу в течение 30 минут. Priority P1 incidents are acknowledged within 30 minutes and resolved within four hours.",
     "Компенсация (Service credit): 10 процентов ежемесячной платы за каждые 0,1 процента ниже порога. Credits are the sole remedy for downtime."],
]

def text_pdf(name, pages, truth_name=None):
    c = canvas.Canvas(str(out / name), pagesize=A4)
    for page in pages:
        y = 790
        for para in page:
            c.setFont("F", 13 if para is page[0] else 11)
            for line in textwrap.wrap(para, 78):
                c.drawString(56, y, line); y -= 17
            y -= 9
        c.showPage()
    c.save()
    (out / "truth" / (name + ".txt")).write_text("\n".join("\n".join(p) for p in pages), encoding="utf-8")

def scanned_pdf(name, pages, size=22):
    font = ImageFont.truetype(FONT, size)
    imgs = []
    for page in pages:
        im = Image.new("L", (1240, 1754), 245)   # 150 dpi A4, paper grey
        d = ImageDraw.Draw(im); y = 110
        for para in page:
            for line in textwrap.wrap(para, 70):
                d.text((100, y), line, fill=20, font=font); y += size + 12
            y += 14
        imgs.append(im.rotate(0.4, fillcolor=245))  # a little skew like a real scan
    imgs[0].save(out / name, save_all=True, append_images=imgs[1:], resolution=150)
    (out / "truth" / (name + ".txt")).write_text("\n".join("\n".join(p) for p in pages), encoding="utf-8")

text_pdf("supply-contract-ru.pdf", RU_PAGES)
text_pdf("nda-en.pdf", EN_PAGES)
text_pdf("sla-bilingual.pdf", MIX_PAGES)
scanned_pdf("scan-contract-ru.pdf", RU_PAGES[:2])
scanned_pdf("scan-nda-en.pdf", EN_PAGES[:1])

# ---- DOCX ----
def docx(name, title, paras, table, bullets):
    d = Document(); d.add_heading(title, 1)
    for t in paras: d.add_paragraph(t)
    for b in bullets: d.add_paragraph(b, style="List Bullet")
    tb = d.add_table(rows=len(table), cols=len(table[0])); tb.style = "Table Grid"
    for i, row in enumerate(table):
        for j, v in enumerate(row): tb.cell(i, j).text = v
    d.save(out / name)
    flat = [title, *paras, *bullets, *[c for r in table for c in r]]
    (out / "truth" / (name + ".txt")).write_text("\n".join(flat), encoding="utf-8")

docx("lease-ru.docx", "Договор аренды нежилого помещения",
     ["Арендодатель передаёт Арендатору во временное пользование помещение площадью 84 квадратных метра по адресу: Санкт-Петербург, Невский проспект, дом 12.",
      "Срок аренды составляет одиннадцать месяцев. Арендная плата вносится ежемесячно не позднее пятого числа."],
     [["Статья расходов", "Сумма, руб."], ["Арендная плата", "145 000"], ["Коммунальные услуги", "18 300"], ["Страховой депозит", "290 000"]],
     ["Арендатор обязан содержать помещение в исправном состоянии.", "Субаренда допускается только с письменного согласия Арендодателя."])
docx("services-en.docx", "Master Services Agreement",
     ["The Supplier shall provide the consulting services described in each Statement of Work with due skill and care.",
      "Invoices are payable within thirty days of receipt. Late payments bear interest at four percent per annum."],
     [["Deliverable", "Fee (USD)"], ["Discovery workshop", "12,500"], ["Implementation", "48,000"], ["Support retainer", "3,200"]],
     ["Either party may terminate for convenience on sixty days notice.", "Liability is capped at the fees paid in the prior twelve months."])

# ---- PPTX ----
def pptx(name, slides):
    pr = Presentation(); truth = []
    for title, body, notes in slides:
        s = pr.slides.add_slide(pr.slide_layouts[1]); s.shapes.title.text = title
        s.placeholders[1].text = "\n".join(body)
        s.notes_slide.notes_text_frame.text = notes
        truth += [title, *body, notes]
    pr.save(out / name); (out / "truth" / (name + ".txt")).write_text("\n".join(truth), encoding="utf-8")

pptx("compliance-training-ru.pptx", [
    ("Обучение по противодействию коррупции", ["Подарки дороже трёх тысяч рублей запрещены", "О конфликте интересов сообщайте в комплаенс-службу"], "Напомнить про горячую линию доверия."),
    ("Персональные данные", ["Согласие субъекта оформляется письменно", "Срок хранения определяется целями обработки"], "Привести пример из практики отдела кадров."),
    ("Ответственность", ["Дисциплинарная и уголовная ответственность", "Порядок проведения служебного расследования"], "Ответить на вопросы аудитории."),
])
pptx("board-update-en.pptx", [
    ("Quarterly Board Update", ["Revenue grew twelve percent year over year", "Two regulatory filings completed on time"], "Mention the audit committee feedback."),
    ("Litigation Overview", ["Three open matters, none material", "Settlement reserve unchanged"], "Counsel is available for questions."),
])

# ---- XLSX ----
wb = openpyxl.Workbook(); ws = wb.active; ws.title = "Платежи"
rows_ru = [["Контрагент", "Договор", "Сумма, руб.", "Дата оплаты"],
           ["ООО «Северный ветер»", "17/2026", 744000, "2026-03-21"],
           ["АО «Балтийская логистика»", "22/2026", 1250000, "2026-04-03"],
           ["ИП Смирнов", "31/2026", 86500, "2026-04-18"]]
for r in rows_ru: ws.append(r)
ws2 = wb.create_sheet("Итоги"); ws2.append(["Показатель", "Значение"]); ws2.append(["Всего оплачено", 2080500]); ws2.append(["Просрочено платежей", 0])
wb.save(out / "payments-ru.xlsx")
(out / "truth" / "payments-ru.xlsx.txt").write_text("\n".join(str(c) for r in [*rows_ru, ["Показатель", "Значение"], ["Всего оплачено", 2080500], ["Просрочено платежей", 0]] for c in r), encoding="utf-8")
wb = openpyxl.Workbook(); ws = wb.active; ws.title = "Invoices"
rows_en = [["Vendor", "Invoice", "Amount USD", "Status"], ["Orion Legal LLP", "INV-1041", 18250, "Paid"], ["Delta Notary", "INV-1057", 640, "Overdue"], ["Kestrel Audit", "INV-1063", 9800, "Pending"]]
for r in rows_en: ws.append(r)
wb.save(out / "invoices-en.xlsx")
(out / "truth" / "invoices-en.xlsx.txt").write_text("\n".join(str(c) for r in rows_en for c in r), encoding="utf-8")

M = [
 ("supply-contract-ru.pdf", "ok", "3", "ДОГОВОР ПОСТАВКИ;;2 480 000;;включая НДС;;Арбитражном суде города Москвы", ""),
 ("nda-en.pdf", "ok", "2", "MUTUAL NON-DISCLOSURE AGREEMENT;;Confidential Information means;;laws of England and Wales", ""),
 ("sla-bilingual.pdf", "ok", "1", "Service Level Agreement;;99.9;;Компенсация", ""),
 ("scan-contract-ru.pdf", "ok", "2", "ДОГОВОР ПОСТАВКИ;;Поставщик обязуется", "ocr"),
 ("scan-nda-en.pdf", "ok", "1", "NON-DISCLOSURE AGREEMENT;;Confidential Information", "ocr"),
 ("lease-ru.docx", "ok", "-", "Договор аренды нежилого помещения;;Невский проспект;;290 000", ""),
 ("services-en.docx", "ok", "-", "Master Services Agreement;;Statement of Work;;12,500", ""),
 ("compliance-training-ru.pptx", "ok", "3", "комплаенс-службу;;Персональные данные;;горячую линию доверия", ""),
 ("board-update-en.pptx", "ok", "2", "Quarterly Board Update;;Settlement reserve unchanged;;audit committee feedback", ""),
 ("payments-ru.xlsx", "ok", "2", "Балтийская логистика;;1250000;;Всего оплачено", ""),
 ("invoices-en.xlsx", "ok", "1", "Orion Legal LLP;;INV-1057;;Overdue", ""),
]
(out / "MANIFEST.tsv").write_text(
    "# Generated legal corpus (RU/EN). file<TAB>expect<TAB>pages<TAB>key phrases;;...<TAB>flags. Truth texts: truth/<file>.txt\n"
    + "\n".join("\t".join(r) for r in M) + "\n", encoding="utf-8")
print("ok", len(M), "files")
